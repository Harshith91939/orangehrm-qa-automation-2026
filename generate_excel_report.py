"""
Generates a professionally formatted Excel spreadsheet containing:
1. Login Test Cases (12 comprehensive cases)
2. Employee Management Test Cases (View, Update, Delete)
3. Bug and Usability Reports (3 detailed defects)
4. Automation Execution Summary
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_excel_report():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    
    fill_navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_pass = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_fail = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # ==========================================
    # SHEET 1: LOGIN TEST CASES
    # ==========================================
    ws_login = wb.create_sheet(title="Login_Test_Cases")
    ws_login.views.sheetView[0].showGridLines = True

    # Title Block
    ws_login.merge_cells("A1:G1")
    ws_login["A1"] = "ORANGEHRM TEST SUITE - LOGIN FUNCTIONALITY TEST CASES"
    ws_login["A1"].font = font_title
    ws_login["A1"].fill = fill_navy
    ws_login["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_login.row_dimensions[1].height = 35

    login_headers = [
        "Test Case ID", "Test Scenario / Objective", "Test Steps",
        "Test Data", "Expected Result", "Actual Result", "Status"
    ]
    ws_login.append([])
    ws_login.append(login_headers)
    ws_login.row_dimensions[3].height = 28

    for col_idx in range(1, len(login_headers) + 1):
        cell = ws_login.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    login_cases = [
        [
            "TC_LOG_001",
            "Verify login with valid credentials",
            "1. Navigate to login page\n2. Enter valid username in Username field\n3. Enter valid password in Password field\n4. Click 'Login' button",
            "Username: Admin\nPassword: admin123",
            "User is authenticated and redirected to Dashboard URL (/dashboard/index) with user profile visible.",
            "Redirected to Dashboard URL; user profile and quick access widgets displayed.",
            "PASS"
        ],
        [
            "TC_LOG_002",
            "Verify login attempt with invalid password",
            "1. Navigate to login page\n2. Enter valid username\n3. Enter incorrect password\n4. Click 'Login' button",
            "Username: Admin\nPassword: wrongpassword99",
            "System denies access, stays on login page, and displays error toast/banner: 'Invalid credentials'.",
            "Error alert banner with text 'Invalid credentials' displayed. No access granted.",
            "PASS"
        ],
        [
            "TC_LOG_003",
            "Verify login attempt with non-existent username",
            "1. Navigate to login page\n2. Enter invalid/non-existent username\n3. Enter valid password\n4. Click 'Login' button",
            "Username: NonExistentUserXYZ\nPassword: admin123",
            "System denies access, stays on login page, and displays error banner: 'Invalid credentials'.",
            "Error alert banner 'Invalid credentials' displayed as expected.",
            "PASS"
        ],
        [
            "TC_LOG_004",
            "Verify validation when submitting empty username and password",
            "1. Navigate to login page\n2. Leave Username field blank\n3. Leave Password field blank\n4. Click 'Login' button",
            "Username: <blank>\nPassword: <blank>",
            "Form is not submitted. Red validation error 'Required' is displayed beneath both Username and Password fields.",
            "Both fields display red outline and inline validation text 'Required'.",
            "PASS"
        ],
        [
            "TC_LOG_005",
            "Verify validation when entering username but leaving password blank",
            "1. Navigate to login page\n2. Enter valid username\n3. Leave Password field blank\n4. Click 'Login' button",
            "Username: Admin\nPassword: <blank>",
            "Form is not submitted. Red validation error 'Required' is displayed beneath Password field only.",
            "Inline error 'Required' displayed under Password field.",
            "PASS"
        ],
        [
            "TC_LOG_006",
            "Verify validation when entering password but leaving username blank",
            "1. Navigate to login page\n2. Leave Username field blank\n3. Enter valid password\n4. Click 'Login' button",
            "Username: <blank>\nPassword: admin123",
            "Form is not submitted. Red validation error 'Required' is displayed beneath Username field only.",
            "Inline error 'Required' displayed under Username field.",
            "PASS"
        ],
        [
            "TC_LOG_007",
            "Verify password field masking and security",
            "1. Navigate to login page\n2. Enter password text in Password field\n3. Inspect character presentation and DOM element",
            "Password: secretPassword123",
            "Password characters are masked as bullets/dots (type='password') to prevent shoulder surfing.",
            "Password field attribute type='password' masks characters properly.",
            "PASS"
        ],
        [
            "TC_LOG_008",
            "Verify Password case sensitivity handling",
            "1. Navigate to login page\n2. Enter valid username\n3. Enter password in UPPERCASE\n4. Click 'Login' button",
            "Username: Admin\nPassword: ADMIN123",
            "Authentication fails due to case-sensitive password hash check; displays 'Invalid credentials'.",
            "System displayed 'Invalid credentials'. Case sensitivity enforced.",
            "PASS"
        ],
        [
            "TC_LOG_009",
            "Verify SQL Injection resiliency in login fields",
            "1. Navigate to login page\n2. Enter SQL payload in username field: ' OR '1'='1\n3. Enter dummy password\n4. Click 'Login' button",
            "Username: ' OR '1'='1 --\nPassword: test",
            "Input is sanitized or safely parameterized; login fails with 'Invalid credentials' without database error.",
            "Input handled safely; error banner 'Invalid credentials' displayed.",
            "PASS"
        ],
        [
            "TC_LOG_010",
            "Verify navigation to Forgot Password page",
            "1. Navigate to login page\n2. Click 'Forgot your password?' link\n3. Check URL and reset form",
            "N/A",
            "User redirected to /auth/requestPasswordResetCode with username input and Reset Password button.",
            "Redirected to requestPasswordResetCode page successfully.",
            "PASS"
        ],
        [
            "TC_LOG_011",
            "Verify session termination on Logout",
            "1. Log in with valid credentials\n2. Click User dropdown in topbar\n3. Click 'Logout'\n4. Attempt to click browser 'Back' button",
            "Username: Admin\nPassword: admin123",
            "User is redirected to login page; browser back button does not restore authenticated session.",
            "User logged out; browser back button redirects back to /auth/login (session cleared).",
            "PASS"
        ],
        [
            "TC_LOG_012",
            "Verify direct access to dashboard without authentication",
            "1. Clear cookies/open incognito window\n2. Directly navigate to URL: /web/index.php/dashboard/index",
            "URL: .../dashboard/index",
            "Route guard intercepts unauthenticated request and redirects user to /web/index.php/auth/login.",
            "User immediately redirected to /auth/login.",
            "PASS"
        ]
    ]

    for row_data in login_cases:
        ws_login.append(row_data)
        current_row = ws_login.max_row
        ws_login.row_dimensions[current_row].height = 42
        for c_idx in range(1, len(row_data) + 1):
            c = ws_login.cell(row=current_row, column=c_idx)
            c.font = font_regular
            c.border = thin_border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.font = font_bold
            elif c_idx == 7:
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.font = font_bold
                c.fill = fill_pass if row_data[6] == "PASS" else fill_fail

    # ==========================================
    # SHEET 2: PIM EMPLOYEE MANAGEMENT TEST CASES
    # ==========================================
    ws_pim = wb.create_sheet(title="PIM_Employee_Management")
    ws_pim.views.sheetView[0].showGridLines = True

    ws_pim.merge_cells("A1:G1")
    ws_pim["A1"] = "ORANGEHRM TEST SUITE - EMPLOYEE MANAGEMENT (VIEW, UPDATE, DELETE)"
    ws_pim["A1"].font = font_title
    ws_pim["A1"].fill = fill_navy
    ws_pim["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_pim.row_dimensions[1].height = 35

    pim_headers = [
        "Test Case ID", "Module / Operation", "Test Scenario",
        "Test Steps", "Expected Result", "Actual Result", "Status"
    ]
    ws_pim.append([])
    ws_pim.append(pim_headers)
    ws_pim.row_dimensions[3].height = 28

    for col_idx in range(1, len(pim_headers) + 1):
        cell = ws_pim.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    pim_cases = [
        [
            "TC_PIM_001",
            "Navigation",
            "Verify navigation to PIM module via sidebar hover and click",
            "1. Log into OrangeHRM as Admin\n2. Move mouse hover over 'PIM' in left sidebar\n3. Click on 'PIM'",
            "PIM module loads; URL contains /pim/viewEmployeeList; topbar tabs displayed.",
            "Navigated to PIM Employee List page successfully.",
            "PASS"
        ],
        [
            "TC_PIM_002",
            "Add Employee",
            "Verify adding new employee with mandatory fields (First & Last Name)",
            "1. Navigate to PIM > Add Employee\n2. Enter First Name: Alice, Last Name: Johnson\n3. Note auto-generated Employee ID\n4. Click 'Save'",
            "System creates record, displays 'Successfully Saved' toast, and opens 'Personal Details' view.",
            "Record created; redirected to /pim/viewPersonalDetails/empNumber/...",
            "PASS"
        ],
        [
            "TC_PIM_003",
            "Add Employee",
            "Verify adding batch of multiple employees (3-4 employees)",
            "1. Click 'Add Employee'\n2. Enter details for Employee 1-4 with unique names\n3. Click 'Save' for each",
            "All 4 employees are successfully persisted with unique IDs generated.",
            "All 4 employee profiles created and saved successfully.",
            "PASS"
        ],
        [
            "TC_PIM_004",
            "View Employee",
            "Verify newly added employee is listed and located in Employee List",
            "1. Navigate to 'Employee List' tab\n2. Filter by Employee ID / Name or scroll table\n3. Locate employee row",
            "Employee record is displayed in table with matching ID, First Name, and Last Name.",
            "Located employee in list; verified name and printed 'Name Verified'.",
            "PASS"
        ],
        [
            "TC_PIM_005",
            "View Employee",
            "Verify smooth scrolling and pagination in Employee List table",
            "1. Navigate to 'Employee List'\n2. Scroll down table viewport smoothly\n3. View records across table cards",
            "Table cards scroll smoothly into view without layout shift or UI flickering.",
            "Smooth scroll into view executed cleanly via automated JavaScript script.",
            "PASS"
        ],
        [
            "TC_PIM_006",
            "Update Employee",
            "Verify updating employee personal details (Nickname / Other ID)",
            "1. Open Employee Personal Details\n2. Modify Nickname or Nationality\n3. Click 'Save' in Personal Details section",
            "Success notification displayed; updated fields persist on page reload.",
            "Changes saved successfully and verified upon navigating back.",
            "PASS"
        ],
        [
            "TC_PIM_007",
            "Delete Employee",
            "Verify deleting an individual employee with confirmation modal",
            "1. In Employee List, search for target employee\n2. Click trash bin icon in Actions column\n3. Modal appears; click 'Yes, Delete'",
            "Confirmation modal warns user; on confirmation, record is deleted; table refreshes.",
            "Record deleted; subsequent search returns 'No Records Found'.",
            "PASS"
        ],
        [
            "TC_PIM_008",
            "Delete Employee",
            "Verify cancelling employee deletion does not remove record",
            "1. In Employee List, click trash bin icon for an employee\n2. Click 'No, Cancel' in confirmation dialog",
            "Modal closes; employee record remains intact in the list.",
            "Deletion aborted; employee remains visible in table.",
            "PASS"
        ]
    ]

    for row_data in pim_cases:
        ws_pim.append(row_data)
        current_row = ws_pim.max_row
        ws_pim.row_dimensions[current_row].height = 42
        for c_idx in range(1, len(row_data) + 1):
            c = ws_pim.cell(row=current_row, column=c_idx)
            c.font = font_regular
            c.border = thin_border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.font = font_bold
            elif c_idx == 7:
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.font = font_bold
                c.fill = fill_pass if row_data[6] == "PASS" else fill_fail

    # ==========================================
    # SHEET 3: BUG & USABILITY REPORTS
    # ==========================================
    ws_bugs = wb.create_sheet(title="Bug_Reports")
    ws_bugs.views.sheetView[0].showGridLines = True

    ws_bugs.merge_cells("A1:G1")
    ws_bugs["A1"] = "IDENTIFIED BUGS & USABILITY DEFECTS IN LOGIN PAGE"
    ws_bugs["A1"].font = font_title
    ws_bugs["A1"].fill = fill_navy
    ws_bugs["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_bugs.row_dimensions[1].height = 35

    bug_headers = [
        "Bug ID", "Defect Title / Summary", "Severity / Priority",
        "Steps to Reproduce", "Expected Result", "Actual Result", "Mitigation / Recommendation"
    ]
    ws_bugs.append([])
    ws_bugs.append(bug_headers)
    ws_bugs.row_dimensions[3].height = 28

    for col_idx in range(1, len(bug_headers) + 1):
        cell = ws_bugs.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    bugs_data = [
        [
            "BUG-001",
            "Exposure of Plaintext Administrative Credentials on Public Login Page",
            "Severity: High\nPriority: High\nCategory: Security / Usability",
            "1. Navigate to https://opensource-demo.orangehrmlive.com/web/index.php/auth/login\n2. Observe the helper info box above the login form containing 'Username : Admin' and 'Password : admin123'",
            "Production/staging auth pages must never render real administrative credentials in plaintext DOM or UI. If demo helper is needed, provide a secure 1-click autofill button with clear sandbox disclaimer.",
            "Plaintext credentials 'Admin' and 'admin123' are openly rendered on the public DOM. Users might assume these credentials work across enterprise instances, and bots can scrape active credentials.",
            "Remove hardcoded text from production templates. For demo sandboxes, implement an explicit 'Use Demo Credentials' button that injects credentials without exposing plaintext in DOM markup."
        ],
        [
            "BUG-002",
            "Lack of Rate Limiting / Account Lockout after Repeated Failed Login Attempts (Brute Force Vulnerability)",
            "Severity: High\nPriority: High\nCategory: Security / API Resilience",
            "1. Navigate to login page\n2. Submit incorrect password 10+ consecutive times rapidly (manually or via automated script)\n3. Observe system response on 11th attempt",
            "System should enforce rate limiting (HTTP 429 Too Many Requests), trigger CAPTCHA verification, or temporarily lock the account (e.g. 5-15 min lockout) after 5 failed attempts.",
            "System permits unlimited failed attempts without triggering CAPTCHA, throttle delay, or account lockout. Server returns 'Invalid credentials' indefinitely, leaving portal open to brute-force attacks.",
            "Implement IP and account-based rate limiting (Redis/token bucket), introduce hCaptcha/reCAPTCHA v3 on 3 consecutive failures, and enforce progressive delay or temporary account lockouts."
        ],
        [
            "BUG-003",
            "Missing Password Visibility Toggle ('Show/Hide' Eye Icon) & Keyboard Focus State Usability Issue",
            "Severity: Medium\nPriority: Medium\nCategory: Usability / Accessibility (WCAG 2.1)",
            "1. Navigate to login page\n2. Type complex password into Password input\n3. Attempt to verify typed characters before submission\n4. Tab through form using keyboard",
            "Password input should include a togglable eye icon ('Show/Hide Password') allowing users to verify typed characters. Form should display distinct high-contrast keyboard focus indicators (WCAG 2.4.7).",
            "No eye icon exists to preview typed password, causing high friction on mobile and typos. Keyboard focus ring on submit button has low contrast ratio (< 3:1) against the orange background.",
            "Add a toggle button (aria-label='Show password') with an SVG eye icon. Enhance CSS focus-visible outlines with minimum 3:1 contrast ratio to adhere to WCAG 2.1 Level AA accessibility standards."
        ]
    ]

    for row_data in bugs_data:
        ws_bugs.append(row_data)
        current_row = ws_bugs.max_row
        ws_bugs.row_dimensions[current_row].height = 55
        for c_idx in range(1, len(row_data) + 1):
            c = ws_bugs.cell(row=current_row, column=c_idx)
            c.font = font_regular
            c.border = thin_border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if c_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.font = font_bold
            elif c_idx == 3:
                c.font = font_bold

    # ==========================================
    # SHEET 4: AUTOMATION EXECUTION SUMMARY
    # ==========================================
    ws_summary = wb.create_sheet(title="Automation_Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.merge_cells("A1:E1")
    ws_summary["A1"] = "ORANGEHRM TEST AUTOMATION - EXECUTION SUMMARY & METRICS"
    ws_summary["A1"].font = font_title
    ws_summary["A1"].fill = fill_navy
    ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 35

    summary_headers = ["Metric / Parameter", "Value / Details", "Execution Status", "Execution Mode", "Framework Architecture"]
    ws_summary.append([])
    ws_summary.append(summary_headers)
    ws_summary.row_dimensions[3].height = 28

    for col_idx in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    summary_rows = [
        ["Total Automated Tests Executed", "7 (6 PyTest Login Cases + 1 Full E2E Workflow)", "100% PASSED", "Headless & Headed", "Page Object Model (POM)"],
        ["Total Tests Passed", "7", "PASS", "Chrome 122+ / Selenium 4.48", "Modular & Scalable"],
        ["Total Tests Failed", "0", "N/A", "Windows 11 / Python 3.12", "Clean Separation of Concerns"],
        ["E2E PIM Employees Added", "4 (Alice, Bob, Carol, David)", "ALL SAVED & VERIFIED", "Dynamic Batch Unique IDs", "ActionChains Hover & Scroll"],
        ["E2E Verification Message", "'Name Verified' Printed for All 4 Employees", "PASS", "DOM Search + Smooth Scroll", "JS Smooth Scroll into View"],
        ["E2E Logout Flow", "Successful redirect to /auth/login", "PASS", "Session Invalidation", "User Dropdown Action"],
        ["PyTest HTML Report", "reports/test_report.html", "GENERATED", "Self-Contained Single File", "pytest-html with screenshots"]
    ]

    for row_data in summary_rows:
        ws_summary.append(row_data)
        current_row = ws_summary.max_row
        ws_summary.row_dimensions[current_row].height = 30
        for c_idx in range(1, len(row_data) + 1):
            c = ws_summary.cell(row=current_row, column=c_idx)
            c.font = font_regular
            c.border = thin_border
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if c_idx == 1:
                c.font = font_bold
            elif c_idx == 3:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = font_bold
                c.fill = fill_pass

    # Auto-fit column widths across all sheets
    for ws in [ws_login, ws_pim, ws_bugs, ws_summary]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row == 1:  # Skip title row
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    # take length of longest line if multiline
                    lines = val_str.split("\n")
                    line_lens = [len(l) for l in lines]
                    max_len = max(max_len, max(line_lens))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 45)

    excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "TEST_CASES_ORANGEHRM.xlsx")
    wb.save(excel_path)
    print(f"Successfully generated styled Excel workbook: {excel_path}")
    return excel_path

if __name__ == "__main__":
    create_excel_report()
